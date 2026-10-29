#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera el entregable de la Práctica: Diseño modular y arquitectura inicial del proyecto
para la materia Desarrollo de Aplicaciones y Servicios Virtuales (Ing. Alejandro Montes Moreno).
Proyecto: COTIZA-ELEV.
"""
import os
import sys
from io import BytesIO

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Importamos datos del proyecto
HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import content_a
import content_b

NAVY = RGBColor(0x1F, 0x38, 0x64)
NAVY_HEX = "1F3864"
ACCENT_HEX = "2E5C9A"
LIGHT_HEX = "EDF2F9"
GREY = RGBColor(0x59, 0x59, 0x59)
BODY_FONT = "Arial"

OUT_DIR = os.path.join(os.path.dirname(HERE), "01_Entregables")
DOCX_OUT = os.path.join(OUT_DIR, "Practica_Diseno_Modular_y_Arquitectura_COTIZA-ELEV.docx")
PDF_OUT = os.path.join(OUT_DIR, "Practica_Diseno_Modular_y_Arquitectura_COTIZA-ELEV.pdf")


# ---------------------------------------------------------------- utilidades XML
def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(cell, top=55, bottom=55, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def para_border(paragraph, position="bottom", color=NAVY_HEX, size=10, space=3):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement("w:" + position)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    borders.append(el)
    pPr.append(borders)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pPr.append(el)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    return run


def set_table_fixed(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    total = int(sum(widths_cm) * 567)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    table._tbl.insert(list(table._tbl).index(tblPr) + 1, grid)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(w * 567)))
            tcW.set(qn("w:type"), "dxa")


# ---------------------------------------------------------------- constructor
class Doc:
    def __init__(self):
        self.d = Document()
        self._setup_styles()
        self._setup_footer()

    def _setup_styles(self):
        st = self.d.styles["Normal"]
        st.font.name = BODY_FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.15
        pf.space_after = Pt(5)
        for s in self.d.sections:
            s.top_margin = Cm(2.0)
            s.bottom_margin = Cm(1.9)
            s.left_margin = Cm(2.0)
            s.right_margin = Cm(2.0)

    def _setup_footer(self):
        sec = self.d.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("COTIZA-ELEV · Práctica: Diseño modular y arquitectura inicial · Página ")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        add_field(p, "PAGE")
        r2 = p.add_run(" de ")
        r2.font.size = Pt(8)
        r2.font.color.rgb = GREY
        add_field(p, "NUMPAGES")
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    def h1(self, texto, numero=None, page_break_before=False):
        etiqueta = ("%s. %s" % (numero, texto)) if numero is not None else texto
        p = self.d.add_paragraph()
        if page_break_before:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(etiqueta)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
        para_border(p, "bottom", NAVY_HEX, 8, 3)
        keep_with_next(p)
        return p

    def h2(self, texto):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(texto)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x2E, 0x5C, 0x9A)
        keep_with_next(p)
        return p

    def h3(self, texto):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(texto)
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        keep_with_next(p)
        return p

    def p(self, texto, align="justify", size=10, italic=False, space_after=5):
        par = self.d.add_paragraph()
        par.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                         "center": WD_ALIGN_PARAGRAPH.CENTER,
                         "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        par.paragraph_format.space_after = Pt(space_after)
        r = par.add_run(texto)
        r.font.size = Pt(size)
        r.font.italic = italic
        return par

    def nota(self, texto):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.left_indent = Cm(0.4)
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(6)
        r = par.add_run(texto)
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = GREY
        return par

    def bullets(self, items):
        for it in items:
            par = self.d.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(2.5)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if isinstance(it, tuple):
                r = par.add_run(it[0] + " ")
                r.font.bold = True
                r.font.size = Pt(9.5)
                r2 = par.add_run(it[1])
                r2.font.size = Pt(9.5)
            else:
                r = par.add_run(it)
                r.font.size = Pt(9.5)

    def table(self, headers, rows, widths, font_size=8.5, header_size=9, zebra=True,
              align_center_cols=()):
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            if i in align_center_cols:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(h)
            r.font.bold = True
            r.font.size = Pt(header_size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade_cell(c, NAVY_HEX)
            cell_margins(c)
        for j, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.05
                if i in align_center_cols:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                partes = str(val).split("\n")
                for k, parte in enumerate(partes):
                    r = par.add_run(parte)
                    r.font.size = Pt(font_size)
                    if i == 0 and len(headers) > 2:
                        r.font.bold = True
                    if k < len(partes) - 1:
                        r.add_break()
                if zebra and j % 2 == 1:
                    shade_cell(cells[i], LIGHT_HEX)
                cell_margins(cells[i])
        set_table_fixed(t, [w * 17.2 / sum(widths) for w in widths])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, buf, caption, width_cm=16.8):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(buf, width=Cm(width_cm))
        par_cap = self.d.add_paragraph()
        par_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par_cap.paragraph_format.space_after = Pt(8)
        r = par_cap.add_run(caption)
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = GREY

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(path)


# ---------------------------------------------------------------- figuras
def generar_diagrama_arquitectura():
    fig, ax = plt.subplots(figsize=(10.8, 6.6))
    ax.set_xlim(0, 10.8)
    ax.set_ylim(0, 6.6)
    ax.axis("off")

    # Título del Diagrama
    ax.text(5.4, 6.35, "ARQUITECTURA MODULAR POR CAPAS — COTIZA-ELEV", ha="center", va="center",
            fontsize=11.5, fontweight="bold", color="#1F3864")
    ax.text(5.4, 6.08, "Flujo de control: Usuarios → Control de Acceso → Servicios de Dominio → Motor de Cálculo → Persistencia & Salidas",
            ha="center", va="center", fontsize=8.2, color="#595959", style="italic")

    # Capa 1: Usuarios / Clientes
    u1 = FancyBboxPatch((0.6, 5.0), 4.6, 0.75, boxstyle="round,pad=0.04,rounding_size=0.08",
                        facecolor="#EDF2F9", edgecolor="#1F3864", linewidth=1.2)
    ax.add_patch(u1)
    ax.text(2.9, 5.48, "Rol Técnico (Campo / Móvil)", ha="center", va="center",
            fontsize=9, fontweight="bold", color="#1F3864")
    ax.text(2.9, 5.20, "Levantamiento de solicitudes, consulta de estado (Sin precios)", ha="center", va="center",
            fontsize=7.5, color="#333333")

    u2 = FancyBboxPatch((5.6, 5.0), 4.6, 0.75, boxstyle="round,pad=0.04,rounding_size=0.08",
                        facecolor="#EDF2F9", edgecolor="#1F3864", linewidth=1.2)
    ax.add_patch(u2)
    ax.text(7.9, 5.48, "Rol Administrador de Ventas (Oficina)", ha="center", va="center",
            fontsize=9, fontweight="bold", color="#1F3864")
    ax.text(7.9, 5.20, "Bandeja, cotizaciones, tarifas, emisión PDF y reportes", ha="center", va="center",
            fontsize=7.5, color="#333333")

    # Flechas a Seguridad
    ax.add_patch(FancyArrowPatch((2.9, 5.0), (2.9, 4.45), arrowstyle="-|>", mutation_scale=10, linewidth=1.2, color="#2E5C9A"))
    ax.add_patch(FancyArrowPatch((7.9, 5.0), (7.9, 4.45), arrowstyle="-|>", mutation_scale=10, linewidth=1.2, color="#2E5C9A"))

    # Capa 2: M01 Seguridad y Enrutamiento
    m01 = FancyBboxPatch((0.6, 3.8), 9.6, 0.65, boxstyle="round,pad=0.04,rounding_size=0.08",
                         facecolor="#1F3864", edgecolor="#1F3864", linewidth=1.2)
    ax.add_patch(m01)
    ax.text(5.4, 4.22, "M01 Acceso y Seguridad (Router, Sesiones, Middleware RBAC, Bloqueo 15 min)",
            ha="center", va="center", fontsize=9.2, fontweight="bold", color="#FFFFFF")
    ax.text(5.4, 3.98, "Valida credenciales, tokens, inactividad (20 min) y filtra rutas por rol antes del controlador",
            ha="center", va="center", fontsize=7.6, color="#D0E0F5")

    # Flecha hacia abajo
    ax.add_patch(FancyArrowPatch((5.4, 3.8), (5.4, 3.35), arrowstyle="-|>", mutation_scale=10, linewidth=1.2, color="#2E5C9A"))

    # Capa 3: Módulos de Aplicación y Negocio
    modulos_boxes = [
        ("M02 Clientes &\nM03 Activos", "Gestión de clientes,\nequipos y validación\nde rangos técnicos", 0.6, 2.15, 2.1, 1.15),
        ("M04 Catálogo &\nTarifas", "Equipos, servicios,\nvigencias y carga\nmasiva CSV", 2.95, 2.15, 2.1, 1.15),
        ("M05 Solicitudes\n(Técnico)", "Requerimientos,\nprevención de\nduplicados y bandeja", 5.3, 2.15, 2.2, 1.15),
        ("M06 Cotizador\nComercial", "Gestión de partidas,\nborradores, duplicado\ny ciclo de estados", 7.75, 2.15, 2.45, 1.15),
    ]
    for tit, dsc, x, y, w, h in modulos_boxes:
        b = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.04,rounding_size=0.08",
                           facecolor="#F4F7FB", edgecolor="#2E5C9A", linewidth=1.1)
        ax.add_patch(b)
        ax.text(x + w / 2, y + h - 0.28, tit, ha="center", va="center", fontsize=8.2, fontweight="bold", color="#1F3864")
        ax.text(x + w / 2, y + 0.32, dsc, ha="center", va="center", fontsize=7.0, color="#444444")

    # Flecha de M06 a M07
    ax.add_patch(FancyArrowPatch((8.98, 2.15), (8.98, 1.62), arrowstyle="-|>", mutation_scale=10, linewidth=1.2, color="#2E5C9A"))

    # Capa 4: Persistencia y Motor de Cálculo
    bd = FancyBboxPatch((0.6, 0.95), 5.8, 0.65, boxstyle="round,pad=0.04,rounding_size=0.08",
                        facecolor="#EDF2F9", edgecolor="#1F3864", linewidth=1.2)
    ax.add_patch(bd)
    ax.text(3.5, 1.37, "Base de Datos MySQL (Capa de Persistencia y Transacciones)", ha="center", va="center",
            fontsize=8.5, fontweight="bold", color="#1F3864")
    ax.text(3.5, 1.12, "Tablas relacionales: usuarios, clientes, activos, tarifas, solicitudes, cotizaciones", ha="center", va="center",
            fontsize=7.2, color="#333333")

    m07 = FancyBboxPatch((6.65, 0.95), 3.55, 0.65, boxstyle="round,pad=0.04,rounding_size=0.08",
                         facecolor="#E3EDF9", edgecolor="#1F3864", linewidth=1.2)
    ax.add_patch(m07)
    ax.text(8.42, 1.37, "M07 Motor de Cálculo (Lógica Pura)", ha="center", va="center",
            fontsize=8.2, fontweight="bold", color="#1F3864")
    ax.text(8.42, 1.12, "Tarifas vigentes, subtotal, IVA, pólizas, correctivos", ha="center", va="center",
            fontsize=7.0, color="#333333")

    # Flechas a Base de datos
    for x_arr in (1.65, 4.0, 6.4):
        ax.add_patch(FancyArrowPatch((x_arr, 2.15), (x_arr, 1.6), arrowstyle="-|>", mutation_scale=8, linewidth=0.9, color="#8FA8C8"))

    # Capa 5: Salidas y Transversal
    out1 = FancyBboxPatch((0.6, 0.12), 4.6, 0.65, boxstyle="round,pad=0.04,rounding_size=0.08",
                          facecolor="#F4F7FB", edgecolor="#2E5C9A", linewidth=1.1)
    ax.add_patch(out1)
    ax.text(2.9, 0.52, "M08 Propuesta PDF y Seguimiento", ha="center", va="center",
            fontsize=8.2, fontweight="bold", color="#1F3864")
    ax.text(2.9, 0.28, "Generador de folios anuales, renderizado PDF y envío SMTP", ha="center", va="center",
            fontsize=7.0, color="#555555")

    out2 = FancyBboxPatch((5.6, 0.12), 4.6, 0.65, boxstyle="round,pad=0.04,rounding_size=0.08",
                          facecolor="#F4F7FB", edgecolor="#2E5C9A", linewidth=1.1)
    ax.add_patch(out2)
    ax.text(7.9, 0.52, "M09 Reportes, Auditoría y Manejo de Errores", ha="center", va="center",
            fontsize=8.2, fontweight="bold", color="#1F3864")
    ax.text(7.9, 0.28, "Dashboard ejecutivo, exportación CSV, auditoría y captura de excepciones", ha="center", va="center",
            fontsize=7.0, color="#555555")

    buf = BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def main():
    print("Iniciando generación de documento de diseño modular y arquitectura...")
    doc = Doc()

    # =========================================================================
    # PORTADA
    # =========================================================================
    p = doc.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD IBEROAMERICANA LEÓN")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(2)

    p = doc.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Desarrollo de Aplicaciones y Servicios Virtuales · Ing. Alejandro Montes Moreno")
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(20)

    p = doc.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PRÁCTICA: DISEÑO MODULAR Y\nARQUITECTURA INICIAL DEL PROYECTO")
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(8)

    p = doc.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COTIZA-ELEV — Sistema Web de Cotizaciones para Elevadores (ELEVESE)")
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x2E, 0x5C, 0x9A)
    para_border(p, "bottom", "C9D6E8", 8, 8)
    p.paragraph_format.space_after = Pt(22)

    doc.table(["Concepto", "Detalle"], [
        ("Proyecto", "COTIZA-ELEV — Plataforma web centralizada de cotizaciones y solicitudes."),
        ("Cliente", "ELEVESE (Instalación, modernización y mantenimiento de elevadores)."),
        ("Asignatura", "Desarrollo de Aplicaciones y Servicios Virtuales."),
        ("Docente titular", "Ing. Alejandro Montes Moreno."),
        ("Equipo de trabajo", "Orlando Villalobos Gutiérrez y Diego Emilio Alférez Vallejo."),
        ("Asignación de roles",
         "• Orlando Villalobos Gutiérrez: Arquitecto / Diseñador de Software y Operador de IA.\n"
         "• Diego Emilio Alférez Vallejo: Validador Técnico y Diseñador."),
        ("Documento base", "Documento de Requerimientos v3.0 (43 RF, 29 RNF, 24 HU, 81 CA, Matriz de Trazabilidad)."),
        ("Fecha de entrega", "Octubre de 2026."),
    ], [4.6, 12.6], font_size=9.5, header_size=9.5)

    p = doc.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    r = p.add_run("León, Guanajuato, México · 2026")
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY
    doc.page_break()

    # =========================================================================
    # 1. OBJETIVO BREVE DEL PROYECTO
    # =========================================================================
    doc.h1("Objetivo breve del proyecto", 1)
    doc.p("COTIZA-ELEV es una plataforma web centralizada diseñada para ELEVESE con el objetivo de estandarizar, "
          "acelerar y controlar el proceso comercial de cotización para equipos de elevación vertical (pasajeros, "
          "carga y montaautos), modernizaciones, pólizas de mantenimiento preventivo y servicios correctivos. "
          "El sistema elimina el uso de hojas de cálculo aisladas y comunicaciones informales por mensajería, "
          "permitiendo que los técnicos en campo levanten solicitudes de refacciones sin exponer montos financieros, "
          "mientras que el administrador de ventas calcula partidas paramétricas con tarifas vigentes y emite propuestas "
          "ejecutivas en PDF con folio correlativo oficial.")
    doc.p("Esta práctica transforma la especificación de requerimientos (verificada previamente en 43 requerimientos "
          "funcionales y 29 no funcionales) en un diseño arquitectónico y modular inicial, delimitando responsabilidades, "
          "interfaces de entrada y salida, relaciones de dependencia y la organización del árbol de archivos antes de "
          "iniciar la fase de codificación.")

    # =========================================================================
    # 2. LISTA DE MÓDULOS PRINCIPALES Y RESPONSABILIDADES
    # =========================================================================
    doc.h1("Lista de módulos principales y responsabilidades", 2)
    doc.p("Para garantizar alta cohesión y bajo acoplamiento, el sistema se estructura en nueve módulos funcionales, "
          "organizados lógicamente en tres capas de servicio:")

    doc.h2("Capa de Datos Base y Administración")
    doc.bullets([
        ("M01 Acceso y seguridad:", "Responsable de la autenticación de usuarios, gestión de credenciales seguras, "
         "control de inactividad (20 minutos), bloqueo temporal por fuerza bruta (15 minutos tras 5 intentos fallidos) "
         "y autorización basada en roles (RBAC) para técnico y administrador de ventas."),
        ("M02 Clientes:", "Responsable del registro, edición, búsqueda predictiva sobre hasta 5,000 registros, "
         "prevención de duplicados por RFC/razón social y baja lógica que preserva el histórico de transacciones."),
        ("M03 Activos (Equipos de elevación):", "Responsable de administrar las fichas técnicas de los elevadores "
         "asociados obligatoriamente a un cliente (capacidad kg, paradas, velocidad, serie, año), validación de rangos "
         "técnicos y consulta del historial de intervenciones adaptado al rol del usuario."),
        ("M04 Catálogo y tarifas:", "Responsable del catálogo maestro de equipos nuevos, refacciones y servicios, "
         "administración de tarifas con rangos de vigencia temporal y carga masiva mediante archivos CSV con reporte de errores."),
    ])

    doc.h2("Capa de Flujo Operativo y Comercial")
    doc.bullets([
        ("M05 Solicitudes (Técnico):", "Responsable de la captura en campo de requerimientos de piezas o servicios por "
         "parte del personal técnico, prevención de solicitudes duplicadas activas, bandeja de pendientes ordenada por antigüedad "
         "y seguimiento de estado sin mostrar información financiera ni precios."),
        ("M06 Cotizador:", "Responsable de orquestar la creación de cotizaciones (desde cero o convertidas desde una solicitud), "
         "gestión dinámica de partidas presupuestales, almacenamiento de borradores, duplicación de cotizaciones previas "
         "y control del ciclo de estados comerciales (borrador, enviada, aceptada, rechazada, vencida)."),
        ("M07 Motor de cálculo:", "Componente desacoplado de lógica pura de negocio. Responsable de procesar reglas paramétricas "
         "de paradas y capacidad, tarifas vigentes, costos por visita en pólizas preventivas, horas/refacciones en correctivos, "
         "descuentos aplicados, base gravable e IVA (16%), garantizando precisión a dos decimales sin intervención manual del precio base."),
    ])

    doc.h2("Capa de Salidas, Control y Soporte")
    doc.bullets([
        ("M08 Propuesta PDF y seguimiento:", "Responsable del renderizado de la propuesta económica formal en formato PDF "
         "bajo plantilla corporativa con membrete, condiciones y firmas; asignación de folio consecutivo anual inmutable; "
         "envío por correo electrónico SMTP y registro del estatus de aceptación o rechazo."),
        ("M09 Reportes, auditoría y errores:", "Responsable de consolidar el tablero ejecutivo de ventas y solicitudes, "
         "exportación tabular en CSV/PDF, registro inmutable de auditoría para operaciones críticas (bitácora) y capturador "
         "centralizado de excepciones con mensajes comprensibles para el usuario."),
    ])

    # =========================================================================
    # 3. TABLA DE MÓDULOS (ENTRADAS, SALIDAS Y REQUERIMIENTOS)
    # =========================================================================
    doc.h1("Tabla de módulos (entradas, salidas y requerimientos)", 3, page_break_before=True)
    doc.p("La siguiente tabla define de manera exhaustiva las responsabilidades, entradas técnicas, salidas producidas "
          "y requerimientos funcionales cubiertos por cada uno de los nueve módulos del sistema:")

    tabla_modulos_data = [
        ("M01 Acceso y seguridad",
         "Autenticación, autorización RBAC, control de sesiones inactivas y bloqueo temporal.",
         "• Correo y contraseña\n• Token de sesión\n• Petición de recuperación\n• Intentos de acceso",
         "• Sesión iniciada\n• Redirección por rol\n• Token temporal\n• Aviso de inactividad / bloqueo",
         "RF-01, RF-02, RF-03,\nRF-04, RF-05, RF-06,\nRF-07"),
        ("M02 Clientes",
         "Registro, edición, búsqueda y baja lógica de clientes con control de duplicados.",
         "• Datos fiscales y contacto\n• Criterio de búsqueda\n• Confirmación de duplicado\n• ID cliente para baja",
         "• Registro de cliente\n• Lista filtrada (≤ 5,000)\n• Alerta de duplicidad\n• Estado inactivo (baja lógica)",
         "RF-08, RF-09, RF-10,\nRF-11, RF-12"),
        ("M03 Activos",
         "Ficha técnica del equipo, validación de rangos mecánicos e historial de servicios.",
         "• ID cliente\n• Datos técnicos del equipo\n• Confirmación de rangos\n• Petición de historial",
         "• Activo registrado\n• Advertencias de rango\n• Historial sin precios (técnico)\n• Historial completo (ventas)",
         "RF-13, RF-14, RF-15"),
        ("M04 Catálogo y tarifas",
         "Catálogo de conceptos, vigencias de precios y carga masiva mediante CSV.",
         "• Conceptos y especificaciones\n• Precios y fechas de vigencia\n• Archivo CSV de importación",
         "• Catálogo actualizado\n• Tarifa activa por fecha\n• Resumen de importación\n• Bitácora de filas rechazadas",
         "RF-16, RF-17, RF-18,\nRF-19"),
        ("M05 Solicitudes",
         "Captura de necesidades en campo por el técnico, detección de duplicados y bandeja de ventas.",
         "• Cliente, activo, pieza, urgencia y nota técnica\n• Justificación de duplicidad\n• ID solicitud y motivo de rechazo",
         "• Solicitud 'Pendiente de cotizar'\n• Aviso de solicitud previa\n• Bandeja ordenada por fecha\n• Estatus cotizada/rechazada",
         "RF-20, RF-21, RF-22,\nRF-23, RF-24"),
        ("M06 Cotizador",
         "Gestión de cotizaciones, partidas dinámicas, borradores, duplicación y estados comerciales.",
         "• ID solicitud o cliente/activo\n• Tipo de cotización\n• Selección de partidas y cantidades\n• Acción comercial (guardar/duplicar)",
         "• Cotización estructurada\n• Partidas vinculadas\n• Cotización borrador recuperable\n• Copia duplicada con precios actuales",
         "RF-25, RF-26, RF-31,\nRF-32, RF-33, RF-34"),
        ("M07 Motor de cálculo",
         "Lógica matemática desacoplada: cálculo de partidas, pólizas, correctivos, IVA y descuentos.",
         "• Parámetros del equipo (paradas, kg)\n• Tarifas vigentes catálogo\n• Periodicidad y visitas de póliza\n• Refacciones, horas y % descuento",
         "• Importes desglosados\n• Subtotal calculado\n• Monto de descuento e IVA (16%)\n• Total final a 2 decimales",
         "RF-27, RF-28, RF-29,\nRF-30"),
        ("M08 Propuesta PDF y seguimiento",
         "Emisión de propuesta ejecutiva PDF con folio anual, envío por correo y seguimiento de cierre.",
         "• ID cotización lista\n• Datos membrete y condiciones\n• Correo electrónico del cliente\n• Registro de aceptación/rechazo",
         "• Documento PDF oficial\n• Folio único anual consecutivo\n• Confirmación de envío por correo\n• Estado comercial actualizado",
         "RF-35, RF-36, RF-37,\nRF-38"),
        ("M09 Reportes, auditoría y errores",
         "Tablero directivo, exportación tabular, log inmutable de auditoría y manejo central de fallas.",
         "• Filtros de fecha, cliente y técnico\n• Eventos de auditoría del sistema\n• Excepciones y errores de ejecución",
         "• Métricas e indicadores KPI\n• Archivo exportado CSV/PDF\n• Bitácora de acciones críticas\n• Mensaje de error amigable",
         "RF-39, RF-40, RF-41,\nRF-42, RF-43"),
    ]

    doc.table(["Módulo", "Responsabilidad principal", "Entradas técnicas", "Salidas producidas", "Requerimientos relacionados"],
              tabla_modulos_data, [3.2, 4.4, 3.4, 3.4, 2.8], font_size=8, header_size=8.5)

    # =========================================================================
    # 4. RELACIÓN REQUERIMIENTO - MÓDULO (MATRIZ COMPLETA)
    # =========================================================================
    doc.h1("Relación requerimiento - módulo (matriz de trazabilidad)", 4, page_break_before=True)
    doc.p("A continuación se presenta la asignación técnica de los 43 requerimientos funcionales del sistema con su "
          "historia de usuario asociada, la prueba técnica sugerida para validar su implementación y el módulo responsable "
          "designado para su construcción:")

    rf_desc_map = {}
    for grp in content_a.GRUPOS_RF:
        for rf_id, desc, prio in grp[1]:
            rf_desc_map[rf_id] = desc

    matriz_data = []
    for row in content_b.TRAZA:
        rf_id, hu, ca, test, mod = row
        nom_mod = content_b.MODULOS[mod]
        desc_corta = rf_desc_map.get(rf_id, "")
        matriz_data.append((rf_id, desc_corta, hu, test, nom_mod))

    doc.table(["ID", "Requerimiento funcional", "Historia", "Prueba técnica sugerida", "Módulo responsable"],
              matriz_data, [1.4, 6.0, 1.4, 5.0, 3.4], font_size=7.5, header_size=8.0, align_center_cols=(0, 2))

    doc.nota("Verificación técnica: Los 43 requerimientos funcionales están debidamente trazados sin omisiones a los nueve módulos propuestos.")

    # =========================================================================
    # 5. DIAGRAMA Y REPRESENTACIÓN DE LA ARQUITECTURA
    # =========================================================================
    doc.h1("Diagrama y representación de la arquitectura", 5, page_break_before=True)
    doc.p("El sistema adopta una **Arquitectura en Capas (Layered Architecture)** con diseño modular desacoplado, "
          "diseñada específicamente para el entorno web de ELEVESE (PHP/MySQL en servidor Apache bajo cPanel). "
          "Esta arquitectura garantiza una estricta separación de responsabilidades, evitando que la lógica de cálculo "
          "se mezcle con la interfaz o con el acceso a la base de datos.")

    doc.h2("Representación textual del flujo de control y datos")
    doc.p("El flujo de ejecución de una petición se procesa en siete etapas estandarizadas:")
    doc.p("1. **Usuario e Interfaz de Entrada:** El usuario (Técnico desde dispositivo móvil o Administrador de ventas "
          "desde oficina) interactúa con la interfaz web responsiva (HTML5 / Bootstrap / JavaScript asíncrono).", space_after=3)
    doc.p("2. **Front Controller y Seguridad (M01):** La petición HTTP llega a un punto de entrada único (`public/index.php`). "
          "El `Router` y el middleware de seguridad verifican la sesión activa, controlan el tiempo de inactividad (20 minutos) "
          "y validan los permisos RBAC. Si un técnico intenta acceder a rutas financieras, la petición es rechazada de inmediato "
          "y registrada en auditoría.", space_after=3)
    doc.p("3. **Controladores de Módulo (M02 a M06):** El controlador correspondiente recibe la solicitud, invoca los "
          "validadores de datos de entrada y coordina las entidades del dominio (clientes, equipos, solicitudes o cotizaciones).", space_after=3)
    doc.p("4. **Lógica de Negocio y Motor de Cálculo (M07):** Cuando la operación involucra importes o partidas comerciales, "
          "el Cotizador delega los datos al `Motor de Cálculo`. Este módulo autónomo procesa las fórmulas paramétricas, aplica tarifas "
          "vigentes del catálogo, calcula pólizas o correctivos, deducciones de descuento e impuestos (IVA 16%), retornando "
          "estructuras matemáticas precisas sin alterar la base de datos directamente.", space_after=3)
    doc.p("5. **Capa de Persistencia (Base de Datos):** Los modelos interactúan con MySQL mediante transacciones seguras (PDO) "
          "para garantizar la consistencia referencial de clientes, activos, folios y partidas.", space_after=3)
    doc.p("6. **Generador de Salidas (M08 / M09):** Dependiendo del caso de uso, el sistema genera la respuesta visual en pantalla, "
          "el archivo PDF con plantilla ejecutiva y folio único anual (M08), o las tablas del dashboard y archivos CSV (M09).", space_after=3)
    doc.p("7. **Manejo Central de Errores y Auditoría (M09):** De forma transversal, cualquier excepción es interceptada por "
          "el manejador central de errores, registrando el evento técnico en el archivo de log y presentando al usuario una "
          "explicación amigable en español. Las transacciones sensibles se escriben en la bitácora inmutable.", space_after=6)

    # Incrustamos la figura generada con matplotlib
    buf_diag = generar_diagrama_arquitectura()
    doc.figure(buf_diag, "Figura 1. Diagrama de arquitectura modular por capas y flujo de interacción de COTIZA-ELEV.")

    # =========================================================================
    # 6. ESTRUCTURA INICIAL DE CARPETAS Y ARCHIVOS
    # =========================================================================
    doc.h1("Estructura inicial de carpetas y archivos", 6, page_break_before=True)
    doc.p("Para materializar el diseño modular en el código fuente, se define una estructura organizada de archivos "
          "siguiendo las mejores prácticas para aplicaciones web estructuradas. Se separa estrictamente la raíz pública web "
          "del código del núcleo y la capa de almacenamiento:")

    arbol_texto = (
        "cotiza-elev/\n"
        "├── public/                         # Raíz pública del servidor web (DocumentRoot)\n"
        "│   ├── index.php                   # Punto de entrada único (Front Controller)\n"
        "│   ├── .htaccess                   # Reglas de reescritura hacia index.php\n"
        "│   ├── assets/\n"
        "│   │   ├── css/\n"
        "│   │   │   ├── bootstrap.min.css   # Estilos base responsivos\n"
        "│   │   │   └── app.css             # Estilos corporativos de ELEVESE (paleta navy)\n"
        "│   │   ├── js/\n"
        "│   │   │   ├── app.js              # Manejo global de alertas y temporizador de sesión\n"
        "│   │   │   ├── cotizador.js        # Lógica interactiva de partidas y recálculo visual\n"
        "│   │   │   └── solicitudes.js      # Validación en cliente de solicitudes técnicas\n"
        "│   │   └── img/\n"
        "│   │       ├── logo_elevese.png    # Logotipo institucional para pantalla y PDF\n"
        "│   │       └── favicon.ico\n"
        "├── config/                         # Archivos de configuración del entorno\n"
        "│   ├── app.php                     # Parámetros del sistema (nombre, URL base, timezone)\n"
        "│   └── database.php                # Credenciales y configuración de conexión PDO\n"
        "├── app/                            # Núcleo de la aplicación y lógica del sistema\n"
        "│   ├── Core/                       # Infraestructura transversal del sistema\n"
        "│   │   ├── Router.php              # Despachador de rutas HTTP\n"
        "│   │   ├── Database.php            # Conexión Singleton PDO transaccional\n"
        "│   │   ├── Session.php             # Control de tiempo de sesión y expiración (20 min)\n"
        "│   │   ├── AuthMiddleware.php      # Control de acceso basado en roles (Técnico / Ventas)\n"
        "│   │   └── ErrorHandler.php        # Captura global de excepciones y mensajes amigables\n"
        "│   ├── Modules/                    # Módulos del sistema organizados por dominio\n"
        "│   │   ├── Auth/                   # M01: Autenticación, bloqueo y recuperación\n"
        "│   │   │   ├── AuthController.php\n"
        "│   │   │   └── UserModel.php\n"
        "│   │   ├── Clientes/               # M02: CRUD clientes, búsqueda predictiva y baja lógica\n"
        "│   │   │   ├── ClienteController.php\n"
        "│   │   │   └── ClienteModel.php\n"
        "│   │   ├── Activos/                # M03: Equipos elevadores y validación de rangos\n"
        "│   │   │   ├── ActivoController.php\n"
        "│   │   │   └── ActivoModel.php\n"
        "│   │   ├── Catalogo/               # M04: Catálogo, tarifas vigentes e importación CSV\n"
        "│   │   │   ├── CatalogoController.php\n"
        "│   │   │   ├── CsvImporter.php\n"
        "│   │   │   └── TarifaModel.php\n"
        "│   │   ├── Solicitudes/            # M05: Solicitudes de campo y bandeja de pendientes\n"
        "│   │   │   ├── SolicitudController.php\n"
        "│   │   │   └── SolicitudModel.php\n"
        "│   │   ├── Cotizador/              # M06: Orquestador de cotizaciones y partidas\n"
        "│   │   │   ├── CotizadorController.php\n"
        "│   │   │   └── CotizacionModel.php\n"
        "│   │   ├── MotorCalculo/           # M07: Lógica pura de cálculo matemático e impuestos\n"
        "│   │   │   ├── CalculatorService.php\n"
        "│   │   │   └── TaxRules.php\n"
        "│   │   ├── Propuestas/             # M08: Renderizado PDF, folios anuales y envío correo\n"
        "│   │   │   ├── PdfGeneratorService.php\n"
        "│   │   │   ├── FolioManager.php\n"
        "│   │   │   └── MailerService.php\n"
        "│   │   └── Reportes/               # M09: Métricas KPI, auditoría y bitácora del sistema\n"
        "│   │       ├── DashboardController.php\n"
        "│   │       ├── AuditLogger.php\n"
        "│   │       └── ExporterService.php\n"
        "│   └── Views/                      # Plantillas HTML estructuradas\n"
        "│       ├── layouts/                # Plantillas maestras (encabezado, barra lateral por rol)\n"
        "│       ├── auth/                   # Vistas de login y recuperación\n"
        "│       ├── clientes/               # Vistas de directorio y formulario de clientes\n"
        "│       ├── activos/                # Ficha técnica e historial del equipo\n"
        "│       ├── catalogo/               # Gestión de conceptos y carga CSV\n"
        "│       ├── solicitudes/            # Captura móvil del técnico y bandeja de ventas\n"
        "│       ├── cotizador/              # Pantalla de cotización con partidas dinámicas\n"
        "│       ├── reportes/               # Tablero de control y filtros de exportación\n"
        "│       └── pdf_templates/          # Plantilla HTML/CSS para compilar la propuesta formal\n"
        "├── storage/                        # Datos temporales, archivos generados y registros\n"
        "│   ├── pdfs/                       # Almacenamiento local de propuestas generadas\n"
        "│   ├── uploads/csv/                # Archivos CSV temporales subidos para procesamiento\n"
        "│   └── logs/\n"
        "│       ├── error.log               # Registro técnico de excepciones internas\n"
        "│       └── audit.log               # Bitácora inmutable de eventos críticos del negocio\n"
        "├── tests/                          # Pruebas automatizadas de lógica crítica\n"
        "│   ├── test_motor_calculo.py       # Casos de prueba para fórmulas, IVA y redondeo\n"
        "│   ├── test_rbac_permisos.py       # Verificación de rechazo de URLs financieras a técnicos\n"
        "│   └── test_csv_parser.py          # Batería de pruebas de importación y errores en CSV\n"
        "├── .gitignore                      # Exclusión de credenciales y archivos temporales\n"
        "└── README.md                       # Documentación técnica de despliegue y configuración"
    )

    doc.p(arbol_texto, align="left", size=8.0)

    doc.h2("Justificación de las carpetas principales")
    doc.bullets([
        ("public/:", "Única carpeta expuesta al servidor web. Al ubicar `index.php` aquí y mantener el resto del código "
         "fuera de la raíz pública, se previene que usuarios maliciosos accedan directamente a archivos PHP, logs o configuraciones."),
        ("config/:", "Aísla las variables del entorno y credenciales de base de datos en archivos independientes, permitiendo "
         "migrar entre entornos de desarrollo y producción en cPanel sin modificar el código de la aplicación."),
        ("app/Core/:", "Centraliza la infraestructura base (enrutador, sesiones, control de inactividad, middleware de roles). "
         "Garantiza que cualquier módulo nuevo herede automáticamente la seguridad del sistema."),
        ("app/Modules/:", "Organización orientada a dominio donde cada módulo cuenta con su propio controlador y modelo. "
         "Permite que dos desarrolladores trabajen en módulos distintos sin generar conflictos en Git."),
        ("storage/:", "Directorio fuera del acceso web público con permisos de escritura restringidos, destinado a PDFs emitidos, "
         "archivos temporales y las bitácoras de auditoría exigidas por RNF-20 y RNF-21."),
        ("tests/:", "Contenedor de scripts automatizados para verificar las reglas críticas del negocio antes del despliegue en servidor."),
    ])

    # =========================================================================
    # 7. JUSTIFICACIÓN DEL DISEÑO ELEGIDO (PREGUNTAS GUÍA)
    # =========================================================================
    doc.h1("Justificación del diseño elegido y respuestas a las preguntas guía", 7, page_break_before=True)
    doc.p("Para fundamentar las decisiones arquitectónicas tomadas, se responde formalmente a las siete preguntas guía "
          "establecidas en la rúbrica de la práctica:")

    preguntas_respuestas = [
        ("1. ¿Qué módulo implementará cada requerimiento?",
         "Todos y cada uno de los 43 requerimientos funcionales están mapeados formalmente en la matriz de trazabilidad (Sección 4). "
         "Cada requerimiento cuenta con un único módulo responsable principal: M01 concentra la seguridad (RF-01 a RF-07), M02 los clientes "
         "(RF-08 a RF-12), M03 los activos (RF-13 a RF-15), M04 el catálogo (RF-16 a RF-19), M05 las solicitudes (RF-20 a RF-24), "
         "M06 la orquestación comercial (RF-25, RF-26, RF-31 a RF-34), M07 el cálculo matemático (RF-27 a RF-30), M08 la emisión documental "
         "(RF-35 a RF-38) y M09 los reportes y control (RF-39 a RF-43)."),

        ("2. ¿Algún módulo tiene demasiadas responsabilidades?",
         "Durante el análisis inicial, el Cotizador corría el riesgo de convertirse en un módulo gigante ('God Object') al pretender gestionar "
         "la interfaz de usuario, validar datos, calcular tarifas complejas de elevadores y renderizar el PDF. "
         "Para solucionar esto, se aplicó el principio de responsabilidad única separándolo en tres módulos independientes: "
         "M06 gestiona la sesión comercial y las partidas; M07 encapsula de forma aislada toda la matemática financiera; y M08 se encarga "
         "exclusivamente de generar el PDF y los folios oficiales. Con esta división, ningún módulo excede sus responsabilidades naturales."),

        ("3. ¿Hay requerimientos que todavía no tienen módulo responsable?",
         "No. La verificación técnica automática realizada sobre el documento base v3.0 confirmó que el 100% de los 43 requerimientos funcionales "
         "y los 29 requerimientos no funcionales tienen un módulo asignado. No existen requerimientos huérfanos ni módulos sin requerimientos asociados."),

        ("4. ¿Qué módulo será más difícil de probar?",
         "El módulo M07 (Motor de cálculo). La dificultad radica en la alta cantidad de combinaciones paramétricas: capacidad de carga, "
         "número de paradas adicionales, fecha de vigencia de la tarifa, prorrateo de visitas en pólizas preventivas, horas de mano de obra "
         "en correctivos y cálculo estricto de descuentos con IVA. Un error de redondeo de centavos altera el valor de la propuesta económica. "
         "Por esta razón, este módulo se construirá como código desacoplado y se probará mediante una batería de pruebas unitarias automatizadas "
         "con casos de cálculo manual previo."),

        ("5. ¿Qué módulos podrían reutilizarse en otros proyectos?",
         "Varios módulos fueron diseñados como componentes reutilizables: "
         "1) M01 (Acceso y seguridad): el middleware de roles y control de inactividad de 20 minutos puede trasladarse a cualquier sistema administrativo. "
         "2) M04 (Importador CSV): el validador de estructura de archivos y reporte de líneas rechazadas es genérico. "
         "3) M07 (Motor de reglas numéricas): su lógica de base gravable, descuentos e IVA es aplicable a cualquier módulo de presupuestos. "
         "4) M08 (Generador de PDF y foliador correlativo): el servicio de renderizado de documentos empresariales y numeración anual. "
         "5) M09 (AuditLogger y ErrorHandler): el capturador global de errores y bitácora inmutable de eventos sensibles."),

        ("6. ¿La IA propuso algo demasiado complejo o fuera del alcance?",
         "Sí. Al solicitarle a la IA la arquitectura inicial, propuso una solución basada en microservicios independientes montados sobre "
         "contenedores Docker, comunicados mediante un broker de mensajería asíncrona RabbitMQ, e incluyó un módulo de machine learning "
         "para predecir el costo de modernizaciones. El equipo rechazó categóricamente esta propuesta: para un equipo de 2 desarrolladores, "
         "un periodo de 8 semanas y una infraestructura en cPanel para ELEVESE, los microservicios y brokers introducían una sobreingeniería "
         "injustificada y un riesgo operacional altísimo. Se optó por una arquitectura monolítica modular en capas, limpia y mantenible."),

        ("7. ¿Qué decisión técnica tomó el equipo y cómo la justificó?",
         "El equipo tomó dos decisiones técnicas fundamentales: "
         "1) Separar el Motor de Cálculo (M07) como servicio puro sin acceso directo a base de datos ni a sesión HTTP: esto permite probarlo "
         "de forma unitaria y reutilizarlo indistintamente en el cotizador interactivo, en la generación de PDFs y en los reportes de ventas. "
         "2) Restricción de visibilidad a nivel de servicio y consulta SQL, no solo en la interfaz: para garantizar que el rol técnico nunca "
         "vea precios ni importes (RF-06), la restricción se aplica en el controlador y en las consultas (`SELECT` que omiten montos cuando "
         "el rol es técnico). Si un usuario técnico inspecciona el tráfico de red o la consola del navegador, la respuesta JSON jamás contiene precios."),
    ]

    for preg, resp in preguntas_respuestas:
        doc.h2(preg)
        doc.p(resp)

    # =========================================================================
    # 8. BITÁCORA DE USO DE LA IA
    # =========================================================================
    doc.h1("Bitácora de uso de la IA", 8, page_break_before=True)
    doc.p("En apego a los lineamientos de la práctica y manteniendo honestidad académica, se registra a continuación "
          "el uso de la inteligencia artificial durante las cuatro fases del diseño modular. La bitácora documenta "
          "los prompts utilizados, las propuestas iniciales del modelo, qué elementos fueron aceptados y qué correcciones "
          "manuales realizó el equipo:")

    bitacora_data = [
        ("Diseño",
         "Actúa como arquitecto de software. Con base en nuestro documento de 43 requerimientos funcionales de COTIZA-ELEV, "
         "propón la arquitectura general del sistema web, los componentes principales y cómo interactúan las peticiones.",
         "Propuso una arquitectura distribuida basada en microservicios con backend en Node.js, frontend en React, un bus "
         "de eventos con RabbitMQ y base de datos distribuida en PostgreSQL y Redis.",
         "Se aceptó la separación conceptual entre autenticación, gestión de entidades base, motor de cálculo y módulo de reportes.",
         "Se rechazó la arquitectura de microservicios y el bus de eventos por sobreingeniería innecesaria. Se adaptó a una "
         "arquitectura monolítica modular en capas sobre PHP y MySQL, acorde a la infraestructura real de cPanel de ELEVESE.",
         "Orlando Villalobos Gutiérrez"),

        ("Módulos",
         "Con base en las 24 historias de usuario y los 43 RF, ayúdanos a identificar los módulos del sistema. Define responsabilidad, "
         "entradas, salidas y requerimientos asociados de cada uno. Evita que un módulo tenga demasiadas funciones.",
         "Identificó inicialmente 6 módulos, agrupando la captura de cotizaciones, el cálculo de fórmulas complejas y la generación "
         "de PDFs dentro de un único 'Módulo de Ventas y Cotización'.",
         "Se aceptaron los módulos de Acceso, Clientes, Activos, Catálogo y Solicitudes.",
         "Se corrigió la sobrecarga del módulo comercial partiéndolo en tres módulos independientes con responsabilidad única: "
         "M06 Cotizador (gestión de partidas), M07 Motor de cálculo (lógica matemática pura) y M08 Propuesta PDF (renderizado y folios).",
         "Diego Emilio Alférez Vallejo"),

        ("Estructura",
         "Diseña una estructura de carpetas y archivos para una aplicación web modular en PHP/MySQL, con Front Controller, "
         "seguridad RBAC, separación de módulos, vistas y manejo de almacenamiento para PDFs y logs.",
         "Generó una estructura de árbol completa con carpeta public/ para el Front Controller, app/Core, app/Modules y storage/.",
         "Se aceptó la distribución de directorios, la ubicación protegida del código fuera de public/ y la carpeta storage para logs.",
         "Se agregaron los módulos específicos de la empresa (Activos, Solicitudes, MotorCalculo), se definieron las plantillas "
         "para el PDF corporativo y se agregó la carpeta tests/ con scripts de prueba unitaria para el motor de cálculo y roles.",
         "Orlando Villalobos Gutiérrez"),

        ("Revisión",
         "Revisa la matriz requerimiento-módulo de nuestros 43 RFs. Comprueba si hay requerimientos huérfanos, dependencias circulares "
         "entre módulos o riesgos de seguridad respecto a que el técnico no debe ver precios.",
         "Confirmó la cobertura de los 43 RFs, pero advirtió que si el módulo M05 Solicitudes consultaba precios para precargar el borrador, "
         "podría exponerse información financiera al técnico por descuido en la API.",
         "Se aceptó la observación crítica de seguridad sobre la API de solicitudes.",
         "Se ajustó la regla de diseño: el técnico solo registra descripción y activo; el precio lo vincula exclusivamente M06 "
         "al momento de convertir la solicitud en cotización, garantizando aislamiento total del rol técnico a nivel de base de datos.",
         "Diego Emilio Alférez Vallejo"),
    ]

    doc.table(["Fase", "Prompt usado", "Resultado de la IA", "Qué aceptamos", "Qué modificamos", "Responsable"],
              bitacora_data, [1.8, 4.2, 3.8, 3.6, 3.8, 2.4], font_size=7.5, header_size=8.0)

    # =========================================================================
    # 9. CHECK LIST PARA LA PRÁCTICA
    # =========================================================================
    doc.h1("Check list para la práctica", 9, page_break_before=True)
    doc.p("Verificación formal de los criterios solicitados por el docente para la entrega de la práctica:")

    checklist_data = [
        ("Identificación adecuada de módulos", "Realizado",
         "Se identificaron 9 módulos con nombres claros y roles definidos, cubriendo el ciclo de vida del negocio sin solapamientos."),
        ("Responsabilidad clara de cada módulo", "Realizado",
         "Cada módulo cuenta con una responsabilidad única delimitada en la Sección 2 y en la tabla de módulos de la Sección 3."),
        ("Relación entre requerimientos y módulos", "Realizado",
         "Se vinculó formalmente el 100% de los 43 requerimientos funcionales con su módulo responsable y prueba técnica en la Sección 4."),
        ("Diagrama o representación de arquitectura", "Realizado",
         "Se construyó el diagrama visual de arquitectura por capas (Figura 1) y se detalló textualmente el flujo de 7 pasos."),
        ("Estructura inicial de carpetas y archivos", "Realizado",
         "Se definió el árbol completo del proyecto con separación de public/, app/Core, app/Modules, storage/ y tests/."),
        ("Justificación del diseño elegido", "Realizado",
         "Se respondieron rigurosamente las 7 preguntas guía en la Sección 7, explicando las decisiones y el descarte de sobreingeniería."),
        ("Bitácora de uso de la IA", "Realizado",
         "Se registraron las cuatro fases de trabajo con IA (Diseño, Módulos, Estructura, Revisión), documentando aportes y correcciones."),
    ]

    doc.table(["Criterio a evaluar", "Estado", "Evidencia y justificación"],
              checklist_data, [4.8, 2.2, 10.2], font_size=8.5, header_size=9.0, align_center_cols=(1,))

    # =========================================================================
    # 10. CONCLUSIÓN Y APORTACIÓN DE LOS INTEGRANTES
    # =========================================================================
    doc.h1("Conclusiones y aportación de los integrantes", 10, page_break_before=True)
    doc.p("El desarrollo de esta práctica permitió establecer con claridad los cimientos arquitectónicos de COTIZA-ELEV antes de "
          "escribir la primera línea de código. Planear la estructura modular permitió identificar riesgos críticos que no eran evidentes "
          "en la etapa de requerimientos, tales como la tendencia del módulo de cotización a concentrar demasiadas funciones, la necesidad "
          "de aislar las reglas matemáticas en un motor de cálculo independiente, y la importancia de proteger los datos financieros "
          "desde el modelo de datos y no solo desde las vistas del navegador.")
    doc.p("El uso de la inteligencia artificial sirvió como un acelerador para proponer estructuras y contrastar alternativas de diseño; "
          "sin embargo, la intervención humana del equipo fue indispensable para aterrizar la arquitectura a la realidad del cliente "
          "ELEVESE, descartando esquemas innecesariamente complejos de microservicios y garantizando que la estructura de archivos responda "
          "a un monolito modular sólido, mantenible y desplegable en cPanel.")

    doc.h2("Evidencias de participación y aportación individual")
    doc.bullets([
        ("Orlando Villalobos Gutiérrez (Arquitecto de Software y Operador de IA):",
         "Diseñó la arquitectura general por capas del sistema, coordinó las consultas estructuradas a la IA en las fases de diseño y "
         "estructura de carpetas, formuló el diagrama de componentes y estructuró la distribución del árbol de archivos en public/, "
         "app/ y storage/."),
        ("Diego Emilio Alférez Vallejo (Validador Técnico y Diseñador):",
         "Realizó la validación cruzada entre los 43 requerimientos funcionales y los módulos asignados, detectó la necesidad de separar el "
         "Motor de Cálculo (M07) del Cotizador (M06), revisó las pruebas técnicas sugeridas y validó el aislamiento de seguridad para que "
         "el técnico no tenga acceso a precios a nivel de servicios."),
    ])

    print("Guardando documento .docx en:", DOCX_OUT)
    doc.save(DOCX_OUT)
    print("Documento .docx guardado exitosamente.")


if __name__ == "__main__":
    main()
